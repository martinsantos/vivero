#!/usr/bin/env python3
"""
Scraper Profundo para Plantas Faitful (plantasfaitful.com.ar)
Genera Excel consolidado con el formato de catalogo_productos_2026.xlsx
"""

import os
import re
import json
import time
import requests
from bs4 import BeautifulSoup
from urllib.parse import urljoin, urlparse, unquote
from pathlib import Path
from datetime import datetime
from xml.etree import ElementTree as ET

# Configuración
BASE_DIR = Path(__file__).parent.parent
OUTPUT_DIR = BASE_DIR / "scraped_data"
IMAGES_DIR = OUTPUT_DIR / "images" / "faitful"
OUTPUT_DIR.mkdir(exist_ok=True)
IMAGES_DIR.mkdir(parents=True, exist_ok=True)

HEADERS = {
    'User-Agent': 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
    'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,image/webp,*/*;q=0.8',
    'Accept-Language': 'es-AR,es;q=0.9,en;q=0.8',
}

BASE_URL = "https://www.plantasfaitful.com.ar"
SITEMAP_URL = f"{BASE_URL}/sitemap.xml"


def get_soup(url, retries=3):
    """Obtiene el BeautifulSoup de una URL con reintentos"""
    for attempt in range(retries):
        try:
            response = requests.get(url, headers=HEADERS, timeout=30)
            response.raise_for_status()
            return BeautifulSoup(response.content, 'html.parser')
        except Exception as e:
            print(f"  ⚠️ Intento {attempt+1}/{retries} fallido para {url}: {e}")
            time.sleep(2)
    return None


def get_xml(url, retries=3):
    """Obtiene el contenido XML de una URL"""
    for attempt in range(retries):
        try:
            response = requests.get(url, headers=HEADERS, timeout=30)
            response.raise_for_status()
            return response.content
        except Exception as e:
            print(f"  ⚠️ Intento {attempt+1}/{retries} fallido para {url}: {e}")
            time.sleep(2)
    return None


def clean_text(text):
    """Limpia texto de espacios extra y caracteres especiales"""
    if not text:
        return ""
    return re.sub(r'\s+', ' ', text.strip())


def download_image(url, product_name, index=0):
    """Descarga una imagen y la guarda con nombre descriptivo"""
    try:
        # Limpiar URL - remover parámetros de redimensionamiento para obtener mejor resolución
        clean_url = re.sub(r'-\d+-\d+\.', '.', url)  # Quitar -480-480. etc
        clean_url = re.sub(r'-\d+x\d+\.', '.', clean_url)  # Quitar -480x480. etc
        
        # Limpiar nombre de archivo
        safe_name = re.sub(r'[^\w\-]', '_', product_name)[:50]
        ext = Path(urlparse(url).path).suffix or '.webp'
        filename = f"{safe_name}_{index}{ext}"
        filepath = IMAGES_DIR / filename
        
        if filepath.exists():
            print(f"    ✓ Imagen ya existe: {filename}")
            return str(filepath), url
        
        response = requests.get(clean_url, headers=HEADERS, timeout=30)
        response.raise_for_status()
        
        with open(filepath, 'wb') as f:
            f.write(response.content)
        
        print(f"    📷 Imagen descargada: {filename}")
        return str(filepath), url
    except Exception as e:
        print(f"    ❌ Error descargando imagen: {e}")
        return None, url


def extract_product_urls_from_sitemap():
    """Extrae todas las URLs de productos desde el sitemap"""
    print("\n📍 Obteniendo URLs desde sitemap...")
    
    xml_content = get_xml(SITEMAP_URL)
    if not xml_content:
        print("❌ No se pudo obtener el sitemap")
        return []
    
    # Parsear XML
    try:
        root = ET.fromstring(xml_content)
    except ET.ParseError as e:
        print(f"❌ Error parseando sitemap: {e}")
        return []
    
    # Namespace del sitemap
    ns = {'sm': 'http://www.sitemaps.org/schemas/sitemap/0.9'}
    
    product_urls = []
    for url_elem in root.findall('.//sm:url', ns):
        loc = url_elem.find('sm:loc', ns)
        if loc is not None:
            url = loc.text
            # Solo URLs de productos
            if '/productos/' in url and '?' not in url:
                product_urls.append(url)
    
    print(f"✅ Encontradas {len(product_urls)} URLs de productos")
    return product_urls


def scrape_product(url):
    """Extrae información completa de un producto de Tiendanube"""
    soup = get_soup(url)
    if not soup:
        return None
    
    product = {
        'ID': '',
        'Nombre': '',
        'Categoría': '',
        'Subcategoría': '',
        'Descripción': '',
        'Cuidados': '',
        'Información Adicional': '',
        'Composición': '',
        'Plagas/Usos': '',
        'Dosis': '',
        'Precio': '',
        'Precio Promocional': '',
        'Variantes': '',
        'Stock': '',
        'SKU': '',
        'URL Imagen': '',
        'Imágenes Adicionales': '',
        'URL Producto': url,
        'Fuente': 'Plantas Faitful'
    }
    
    # === NOMBRE ===
    title = soup.find('h1')
    if title:
        product['Nombre'] = clean_text(title.get_text())
    
    if not product['Nombre']:
        # Extraer del URL
        product['Nombre'] = unquote(url.split('/')[-2].replace('-', ' ').title())
    
    # === CATEGORÍAS (Breadcrumbs) ===
    # En Tiendanube/Faitful la clase 'crumb' está en el tag 'a' mismo: <a class="crumb" ...>
    breadcrumbs = soup.select('a.crumb, .breadcrumb a, nav[aria-label="breadcrumb"] a')
    if breadcrumbs:
        crumb_texts = [clean_text(b.get_text()) for b in breadcrumbs if b.get_text().strip()]
        # Filtrar "Inicio" y el nombre del producto (aunque a.crumb usualmente no incluye el producto)
        crumb_texts = [c for c in crumb_texts if c.lower() not in ['inicio', 'home', '']]
        if len(crumb_texts) >= 1:
            product['Categoría'] = crumb_texts[0]
        if len(crumb_texts) >= 2:
            product['Subcategoría'] = crumb_texts[1]
    
    # Si no hay breadcrumbs vía CSS, extraer de JSON-LD
    if not product['Categoría']:
        script_ld = soup.find_all('script', type='application/ld+json')
        for script in script_ld:
            try:
                ld_data = json.loads(script.string)
                # Tiendanube a veces pone el breadcrumb dentro de un objeto principal
                breadcrumb_obj = None
                if isinstance(ld_data, dict):
                    if ld_data.get('@type') == 'BreadcrumbList':
                        breadcrumb_obj = ld_data
                    elif 'breadcrumb' in ld_data and ld_data['breadcrumb'].get('@type') == 'BreadcrumbList':
                        breadcrumb_obj = ld_data['breadcrumb']
                
                if breadcrumb_obj:
                    items = breadcrumb_obj.get('itemListElement', [])
                    # item[0] es Inicio, item[1] es Categoría, item[2] es Subcategoría
                    useful_items = [i.get('item', {}).get('name', i.get('name', '')) for i in items]
                    useful_items = [c for c in useful_items if c.lower() not in ['inicio', 'home', '']]
                    if len(useful_items) >= 1:
                        product['Categoría'] = useful_items[0]
                    if len(useful_items) >= 2:
                        product['Subcategoría'] = useful_items[1]
                    if product['Categoría']:
                        break
            except:
                continue
    
    # === DESCRIPCIÓN Y CUIDADOS ===
    # Tiendanube usa .user-content para descripción
    user_content = soup.select('.user-content, .product-description, .description')
    for content in user_content:
        text = clean_text(content.get_text())
        if text:
            product['Descripción'] = text
            break
    
    # Buscar secciones específicas
    all_text = soup.get_text()
    
    # Cuidados
    cuidados_match = re.search(r'Cuidados[:\s]*(.+?)(?=Descripción|Información|$)', all_text, re.IGNORECASE | re.DOTALL)
    if cuidados_match:
        product['Cuidados'] = clean_text(cuidados_match.group(1)[:500])
    
    # Información adicional
    info_match = re.search(r'Información adicional[:\s]*(.+?)(?=Cuidados|Descripción|$)', all_text, re.IGNORECASE | re.DOTALL)
    if info_match:
        product['Información Adicional'] = clean_text(info_match.group(1)[:500])
    
    # === PRECIO ===
    # Precio tachado (original)
    compare_price = soup.select_one('.js-compare-price-display, .compare-at-price, .price-compare-at, del .money, .was-price')
    if compare_price:
        product['Precio'] = clean_text(compare_price.get_text())
    
    # Precio actual/promocional
    price = soup.select_one('.js-price-display, .product-price, .price .money:not(del .money), .current-price')
    if price:
        price_text = clean_text(price.get_text())
        if product['Precio']:
            product['Precio Promocional'] = price_text
        else:
            product['Precio'] = price_text
    
    # === VARIANTES ===
    variants = []
    variant_selectors = soup.select('.js-variation-option, .variant-option, [data-variant], .product-variant select option')
    for v in variant_selectors:
        variant_text = clean_text(v.get_text())
        if variant_text and variant_text.lower() not in ['seleccionar', 'elegir', '-']:
            variants.append(variant_text)
    
    # También buscar en labels
    variant_labels = soup.select('.js-label-variant, .variant-label, [for*="variant"]')
    for label in variant_labels:
        label_text = clean_text(label.get_text())
        if label_text:
            variants.append(label_text)
    
    product['Variantes'] = ', '.join(set(variants)) if variants else ''
    
    # === SKU ===
    sku = soup.select_one('.sku, [data-sku], .product-sku')
    if sku:
        product['SKU'] = clean_text(sku.get_text())
    
    # === STOCK ===
    stock = soup.select_one('.stock, [data-stock], .product-stock, .availability')
    if stock:
        product['Stock'] = clean_text(stock.get_text())
    
    # === IMÁGENES ===
    images = []
    
    # Imágenes del producto (Tiendanube usa clase específica)
    img_selectors = [
        '.js-product-slide-img',
        '.product-image img',
        '.js-product-thumb',
        '.product-gallery img',
        '[data-zoom-image]',
        '.swiper-slide img'
    ]
    
    for selector in img_selectors:
        imgs = soup.select(selector)
        for img in imgs:
            src = img.get('src') or img.get('data-src') or img.get('data-zoom-image') or img.get('data-srcset', '').split()[0]
            if src and 'mitiendanube.com' in src:
                # Obtener imagen en máxima resolución
                high_res = re.sub(r'-\d+-\d+\.', '-1024-1024.', src)
                if high_res not in images:
                    images.append(high_res)
    
    # También buscar en meta tags og:image
    og_image = soup.find('meta', property='og:image')
    if og_image and og_image.get('content'):
        img_url = og_image['content']
        if img_url not in images:
            images.insert(0, img_url)
    
    if images:
        product['URL Imagen'] = images[0]
        if len(images) > 1:
            product['Imágenes Adicionales'] = ' | '.join(images[1:5])  # Máximo 4 adicionales
    
    # === ID (generar desde nombre/url) ===
    product['ID'] = abs(hash(url)) % 1000000
    
    return product


def generate_excel(products):
    """Genera archivo Excel en el formato de catalogo_productos_2026.xlsx"""
    print("\n" + "="*60)
    print("📊 GENERANDO EXCEL")
    print("="*60)
    
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("⚠️ Instalando openpyxl...")
        os.system("pip install openpyxl")
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Catálogo Faitful"
    
    # Estilos
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="2D5A3D", end_color="2D5A3D", fill_type="solid")
    alt_fill = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Columnas (basadas en catalogo_productos_2026.xlsx)
    columns = [
        'ID', 'Nombre', 'Categoría', 'Descripción', 'Composición', 
        'Plagas/Usos', 'Dosis', 'URL Imagen'
    ]
    
    # Headers
    for col, header in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border
    
    # Datos
    for row, product in enumerate(products, 2):
        # Construir descripción completa
        desc_parts = []
        if product.get('Descripción'):
            desc_parts.append(product['Descripción'])
        if product.get('Cuidados'):
            desc_parts.append(f"CUIDADOS: {product['Cuidados']}")
        if product.get('Información Adicional'):
            desc_parts.append(f"INFO ADICIONAL: {product['Información Adicional']}")
        
        full_desc = ' | '.join(desc_parts) if desc_parts else ''
        
        # Construir categoría con subcategoría
        cat = product.get('Categoría', '')
        subcat = product.get('Subcategoría', '')
        full_cat = f"{cat} > {subcat}" if subcat else cat
        
        data = [
            product.get('ID', ''),
            product.get('Nombre', ''),
            full_cat,
            full_desc,
            product.get('Composición', ''),
            product.get('Plagas/Usos', product.get('Variantes', '')),
            product.get('Dosis', ''),
            product.get('URL Imagen', '')
        ]
        
        for col, value in enumerate(data, 1):
            cell = ws.cell(row=row, column=col, value=str(value) if value else '')
            cell.border = border
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if row % 2 == 0:
                cell.fill = alt_fill
    
    # Ajustar anchos de columna
    column_widths = [10, 40, 30, 60, 20, 20, 15, 60]
    for col, width in enumerate(column_widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    
    # === HOJA DETALLADA ===
    ws_detail = wb.create_sheet("Detalle Completo")
    
    detail_columns = [
        'ID', 'Nombre', 'Categoría', 'Subcategoría', 'Descripción', 
        'Cuidados', 'Info Adicional', 'Precio', 'Precio Promo', 
        'Variantes', 'SKU', 'Stock', 'URL Imagen', 'Imgs Adicionales', 'URL Producto', 'Fuente'
    ]
    
    # Headers detalle
    for col, header in enumerate(detail_columns, 1):
        cell = ws_detail.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border
    
    # Datos detalle
    for row, product in enumerate(products, 2):
        data = [
            product.get('ID', ''),
            product.get('Nombre', ''),
            product.get('Categoría', ''),
            product.get('Subcategoría', ''),
            product.get('Descripción', ''),
            product.get('Cuidados', ''),
            product.get('Información Adicional', ''),
            product.get('Precio', ''),
            product.get('Precio Promocional', ''),
            product.get('Variantes', ''),
            product.get('SKU', ''),
            product.get('Stock', ''),
            product.get('URL Imagen', ''),
            product.get('Imágenes Adicionales', ''),
            product.get('URL Producto', ''),
            product.get('Fuente', '')
        ]
        
        for col, value in enumerate(data, 1):
            cell = ws_detail.cell(row=row, column=col, value=str(value) if value else '')
            cell.border = border
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if row % 2 == 0:
                cell.fill = alt_fill
    
    # Ajustar anchos
    for ws_current in [ws, ws_detail]:
        for col in range(1, 20):
            try:
                ws_current.column_dimensions[get_column_letter(col)].width = 20
            except:
                pass
    
    # Guardar
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = OUTPUT_DIR / f"catalogo_faitful_{timestamp}.xlsx"
    wb.save(filename)
    
    print(f"\n✅ Excel generado: {filename}")
    print(f"   📄 Productos: {len(products)}")
    
    # También copiar al archivo principal
    main_file = BASE_DIR / "catalogo_productos_2026.xlsx"
    try:
        wb.save(main_file)
        print(f"   📄 Actualizado: {main_file}")
    except Exception as e:
        print(f"   ⚠️ No se pudo actualizar archivo principal: {e}")
    
    return str(filename)


def main():
    print("\n" + "🌿"*30)
    print("    SCRAPER PROFUNDO - PLANTAS FAITFUL")
    print("    plantasfaitful.com.ar")
    print("🌿"*30)
    
    start_time = time.time()
    
    # Obtener URLs de productos
    product_urls = extract_product_urls_from_sitemap()
    
    if not product_urls:
        print("❌ No se encontraron productos")
        return
    
    # Scraping de cada producto
    products = []
    total = len(product_urls)
    
    for i, url in enumerate(product_urls, 1):
        print(f"\n[{i}/{total}] Procesando: {url.split('/')[-2]}")
        
        product = scrape_product(url)
        if product and product.get('Nombre'):
            products.append(product)
            print(f"  ✅ {product['Nombre']}")
        else:
            print(f"  ⚠️ Sin datos")
        
        # Pequeña pausa para no sobrecargar
        time.sleep(0.5)
    
    print(f"\n✅ Total productos extraídos: {len(products)}")
    
    # Guardar JSON de respaldo
    json_file = OUTPUT_DIR / "faitful_productos_raw.json"
    with open(json_file, 'w', encoding='utf-8') as f:
        json.dump({
            'productos': products,
            'total': len(products),
            'timestamp': datetime.now().isoformat()
        }, f, ensure_ascii=False, indent=2)
    print(f"💾 JSON guardado: {json_file}")
    
    # Generar Excel
    excel_file = generate_excel(products)
    
    elapsed = time.time() - start_time
    print(f"\n⏱️ Tiempo total: {elapsed/60:.1f} minutos")
    print(f"📁 Archivos generados en: {OUTPUT_DIR}")
    
    return excel_file


if __name__ == "__main__":
    main()
